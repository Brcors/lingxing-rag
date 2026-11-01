"""Text extraction for the document formats accepted by the knowledge index."""
import csv
import io
import json
import re
import tempfile
import threading
from datetime import date, datetime, time
from pathlib import Path


SUPPORTED_SUFFIXES = {".md", ".txt", ".pdf", ".docx", ".xlsx", ".xlsm", ".xls", ".csv", ".tsv"}


def _clean_cell(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        value = value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    value = str(value).replace("\r", " ").replace("\n", "<br>").strip()
    return value.replace("|", "\\|")


def _markdown_table(rows):
    rows = [[_clean_cell(value) for value in row] for row in rows]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""
    width = max(map(len, rows))
    header = (rows[0] + [""] * width)[:width]
    header = [value or f"列 {index + 1}" for index, value in enumerate(header)]
    lines = ["| " + " | ".join(header) + " |",
             "| " + " | ".join("---" for _ in header) + " |"]
    for row in rows[1:]:
        cells = (row + [""] * width)[:width]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def _text_from_bytes(raw):
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    try:
        from charset_normalizer import from_bytes
        match = from_bytes(raw).best()
        if match is not None:
            return str(match)
    except ImportError:
        pass
    return raw.decode("utf-8", errors="replace")


def _extract_docx(path):
    from docx import Document as WordDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = WordDocument(str(path))
    blocks = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            style = (item.style.name or "") if item.style else ""
            heading = re.match(r"Heading\s+(\d+)", style, re.I)
            if heading:
                text = "#" * min(6, int(heading.group(1)) + 1) + " " + text
            blocks.append(text)
        elif isinstance(item, Table):
            table = _markdown_table([[cell.text for cell in row.cells] for row in item.rows])
            if table:
                blocks.append(table)
    return "\n\n".join(blocks)


def _extract_excel(path):
    if path.suffix.lower() == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(path, on_demand=True)
        try:
            sheets = []
            for sheet in workbook.sheets():
                rows = []
                for row_index in range(sheet.nrows):
                    row = []
                    for col_index in range(sheet.ncols):
                        cell = sheet.cell(row_index, col_index)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate.xldate_as_datetime(value, workbook.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        row.append(value)
                    rows.append(row)
                sheets.append((sheet.name, rows))
            return sheets
        finally:
            workbook.release_resources()

    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return [(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
                for sheet in workbook.worksheets]
    finally:
        workbook.close()


class PaddleOCRExtractor:
    """Lazy, single-instance Chinese PaddleOCR fallback for image-only PDF pages."""

    def __init__(self):
        self._model = None
        self._lock = threading.Lock()

    def _get_model(self):
        if self._model is None:
            with self._lock:
                if self._model is None:
                    try:
                        from paddleocr import PaddleOCR
                    except ImportError as error:
                        raise RuntimeError(
                            "该 PDF 页面是扫描图像；请按部署文档安装 PaddleOCR："
                            "python -m pip install -r requirements-ocr.txt"
                        ) from error
                    self._model = PaddleOCR(
                        lang="ch", device="cpu", engine="paddle",
                        use_doc_orientation_classify=False,
                        use_doc_unwarping=False,
                        use_textline_orientation=False,
                    )
        return self._model

    def extract(self, image_path):
        output = self._get_model().predict(str(image_path))
        text = []
        for item in output:
            payload = getattr(item, "json", item)
            if callable(payload):
                payload = payload()
            if isinstance(payload, str):
                payload = json.loads(payload)
            if isinstance(payload, dict):
                payload = payload.get("res", payload)
                if isinstance(payload, dict):
                    text.extend(str(value).strip() for value in payload.get("rec_texts", []) if str(value).strip())
        return "\n".join(text)


def _extract_pdf(path, ocr):
    from pypdf import PdfReader

    pages = []
    reader = PdfReader(str(path), strict=False)
    renderer = None
    try:
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text(extraction_mode="layout") or "").strip()
            # Sparse text layers commonly indicate scanned pages. OCR only those pages.
            visible_characters = len(re.sub(r"\s+", "", text))
            try:
                has_images = bool(page.images)
            except Exception:
                has_images = False
            if visible_characters < 40 and has_images:
                if renderer is None:
                    try:
                        import pypdfium2 as pdfium
                    except ImportError as error:
                        raise RuntimeError(
                            "扫描 PDF 需要 pypdfium2；请按部署文档安装 PaddleOCR 扩展依赖。"
                        ) from error
                    renderer = pdfium.PdfDocument(str(path))
                with tempfile.TemporaryDirectory(prefix="lingxing-rag-pdf-") as temp_dir:
                    image_path = Path(temp_dir) / f"page-{page_number:05d}.png"
                    pdf_page = renderer[page_number - 1]
                    try:
                        pdf_page.render(scale=220 / 72).to_pil().save(image_path)
                    finally:
                        pdf_page.close()
                    recognized = ocr.extract(image_path)
                if recognized:
                    text = text + "\n" + recognized if text else recognized
            if text:
                pages.append(f"## 第 {page_number} 页\n\n{text}")
    finally:
        if renderer is not None:
            renderer.close()
    return "\n\n".join(pages)


class ContentExtractor:
    def __init__(self):
        self.ocr = PaddleOCRExtractor()

    def extract(self, path):
        suffix = path.suffix.lower()
        if suffix == ".md":
            return _text_from_bytes(path.read_bytes())
        if suffix == ".txt":
            return _text_from_bytes(path.read_bytes())
        if suffix == ".pdf":
            return _extract_pdf(path, self.ocr)
        if suffix == ".docx":
            return _extract_docx(path)
        if suffix in {".xlsx", ".xlsm", ".xls"}:
            sheets = _extract_excel(path)
            blocks = []
            for name, rows in sheets:
                table = _markdown_table(rows)
                if table:
                    blocks.append(f"## 工作表：{name}\n\n{table}")
            return "\n\n".join(blocks)
        if suffix in {".csv", ".tsv"}:
            content = _text_from_bytes(path.read_bytes())
            delimiter = "\t" if suffix == ".tsv" else ","
            if suffix == ".csv":
                try:
                    delimiter = csv.Sniffer().sniff(content[:8192], delimiters=",;\t|").delimiter
                except csv.Error:
                    pass
            return _markdown_table(list(csv.reader(io.StringIO(content), delimiter=delimiter)))
        raise ValueError(f"暂不支持的知识库文件格式：{suffix or '(无扩展名)'}")
